from __future__ import annotations

from pathlib import Path
from docx import Document

from app.services.pdf_parser import TextExtractionError, clean_text


def extract_docx_text(path: str | Path) -> str:
    try:
        document = Document(path)
    except Exception as exc:
        raise TextExtractionError("The DOCX file could not be opened or read") from exc

    parts: list[str] = []

    # 1. Headers & Footers (frequently contain candidate contact info, name, title)
    for section in document.sections:
        try:
            if section.header:
                for p in section.header.paragraphs:
                    txt = p.text.strip()
                    if txt and txt not in parts:
                        parts.append(txt)
            if section.footer:
                for p in section.footer.paragraphs:
                    txt = p.text.strip()
                    if txt and txt not in parts:
                        parts.append(txt)
        except Exception:
            pass

    # 2. Main document body paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    # 3. Document tables
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                row_str = " | ".join(c for c in cells if c)
                if row_str and row_str not in parts:
                    parts.append(row_str)

    # 4. Deep XML extraction fallback (captures text inside TextBoxes, Shapes, Frames, Callouts)
    try:
        xml_text_nodes = document.element.xpath("//w:t")
        for node in xml_text_nodes:
            if node.text:
                txt = node.text.strip()
                if txt and txt not in parts and len(txt) > 2:
                    parts.append(txt)
    except Exception:
        pass

    text = clean_text("\n\n".join(parts))

    if not text:
        raise TextExtractionError(
            "No extractable text was found in this DOCX file. If your resume contains images or custom graphical text boxes, please save or export it as a standard text-based PDF or DOCX file."
        )

    return text
